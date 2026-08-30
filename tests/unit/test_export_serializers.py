"""app/modules/exports/serializers.py e puro (recebe DomainSpec + linhas ja
no formato canonico, devolve bytes) — sem tocar banco/request."""

import csv
import io
import json
from datetime import date

import openpyxl

from app.modules.exports.serializers import (
    build_filename,
    to_csv_bytes,
    to_json_bytes,
    to_xlsx_bytes,
)
from app.modules.imports.domains import FINANCIAL_PERIODS, PRODUCTS

ROWS = [
    {"period": date(2026, 1, 1), "revenue": 10000.0, "costs": 3000.0, "expenses": None,
     "marketing_investment": None, "taxes": None},
    {"period": date(2026, 2, 1), "revenue": 12000.0, "costs": 3500.0, "expenses": 1200.0,
     "marketing_investment": 2000.0, "taxes": 500.0},
]


class TestToCsvBytes:
    def test_header_uses_display_names(self):
        content = to_csv_bytes(FINANCIAL_PERIODS, ROWS)
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")))
        header = next(reader)
        assert header == [f.display_name for f in FINANCIAL_PERIODS.fields]

    def test_date_is_iso_and_none_is_empty_string(self):
        content = to_csv_bytes(FINANCIAL_PERIODS, ROWS)
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        rows = list(reader)
        assert rows[0]["Período"] == "2026-01-01"
        assert rows[0]["Despesas"] == ""

    def test_has_utf8_bom(self):
        content = to_csv_bytes(FINANCIAL_PERIODS, ROWS)
        assert content.startswith(b"\xef\xbb\xbf")

    def test_formula_injection_prefix_is_neutralized(self):
        malicious_rows = [{**ROWS[0], "revenue": "=cmd|'/C calc'!A1"}]
        content = to_csv_bytes(FINANCIAL_PERIODS, malicious_rows)
        assert "'=cmd" in content.decode("utf-8-sig")

    def test_boolean_is_localized_not_python_repr(self):
        # "True"/"False" seria um vazamento de inglês num arquivo cujo
        # cabeçalho já é todo em português.
        rows = [{"name": "Produto A", "category": None, "price": None, "cost": None, "active": True}]
        content = to_csv_bytes(PRODUCTS, rows)
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
        assert next(reader)["Ativo"] == "Sim"


class TestToXlsxBytes:
    def test_reads_back_headers_and_values(self):
        content = to_xlsx_bytes(FINANCIAL_PERIODS, ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        sheet = wb.active
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        assert header == [f.display_name for f in FINANCIAL_PERIODS.fields]

        first_data_row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
        # openpyxl sempre devolve datetime ao ler uma celula de data (Excel
        # nao distingue date/datetime no armazenamento) — o que importa aqui
        # e que a celula ficou como tipo de data nativo, nao uma string ISO.
        assert first_data_row[0].date() == date(2026, 1, 1)
        assert first_data_row[1] == 10000.0

    def test_sheet_title_is_domain_display_name(self):
        content = to_xlsx_bytes(FINANCIAL_PERIODS, ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        assert wb.active.title == "Financeiro"


class TestToJsonBytes:
    def test_uses_canonical_field_names_as_keys(self):
        content = to_json_bytes(FINANCIAL_PERIODS, ROWS)
        payload = json.loads(content)
        assert payload[0]["period"] == "2026-01-01"
        assert payload[0]["revenue"] == 10000.0
        assert payload[0]["expenses"] is None

    def test_round_trips_all_rows(self):
        content = to_json_bytes(FINANCIAL_PERIODS, ROWS)
        payload = json.loads(content)
        assert len(payload) == 2
        assert payload[1]["period"] == "2026-02-01"


class TestBuildFilename:
    def test_without_date_range(self):
        assert build_filename("financial_periods", "csv", None, None) == "financial_periods_export.csv"

    def test_with_date_range(self):
        name = build_filename("financial_periods", "xlsx", date(2026, 1, 1), date(2026, 12, 31))
        assert name == "financial_periods_2026-01-01_2026-12-31.xlsx"
