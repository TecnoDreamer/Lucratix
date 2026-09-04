"""Exportação de dados — corte vertical, domínio Financeiro (2026-08-29).
Mesmo nível de permissão de GET /financials (qualquer membership ativa,
inclusive VIEWER) — exportar é ler o mesmo dado que já aparece na tela."""

import csv
import io
import json
from datetime import UTC, datetime

import openpyxl


def _register_login_and_create_org(client, email: str, org_name: str) -> str:
    client.post(
        "/api/v1/auth/register",
        json={"name": "Ana", "email": email, "password": "SenhaForte123!"},
    )
    client.post("/api/v1/auth/login", json={"email": email, "password": "SenhaForte123!"})
    org = client.post("/api/v1/organizations", json={"name": org_name}).json()
    return org["id"]


def _register_and_login(client, name: str, email: str, password: str = "SenhaForte123!") -> None:
    client.post("/api/v1/auth/register", json={"name": name, "email": email, "password": password})
    client.post("/api/v1/auth/login", json={"email": email, "password": password})


def _seed_financial_periods(client, org_id: str) -> None:
    client.put(
        f"/api/v1/organizations/{org_id}/financials",
        json={"period": "2026-01-01", "revenue": 10000, "costs": 3000},
    )
    client.put(
        f"/api/v1/organizations/{org_id}/financials",
        json={"period": "2026-02-01", "revenue": 12000, "costs": 3500},
    )


def test_export_csv_contains_seeded_periods(client):
    org_id = _register_login_and_create_org(client, "exp.csv@example.com", "Empresa Exp CSV")
    _seed_financial_periods(client, org_id)

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "csv"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert 'attachment; filename="financial_periods_export.csv"' in resp.headers["content-disposition"]

    rows = list(csv.DictReader(io.StringIO(resp.content.decode("utf-8-sig"))))
    assert len(rows) == 2
    assert rows[0]["Período"] == "2026-01-01"
    assert rows[0]["Receita"] == "10000.0"


def test_export_exposes_content_disposition_via_cors(client):
    """Sem `expose_headers`, o browser esconde Content-Disposition de um
    `fetch` cross-origin (dev local: frontend em :3000 chamando a API em
    :8000 direto) — `downloadFile` (apps/web/src/lib/api.ts) cairia sempre no
    nome de arquivo genérico "export". Verificado ao vivo antes desta
    correção; nunca deve regredir silenciosamente."""
    org_id = _register_login_and_create_org(client, "exp.cors@example.com", "Empresa Exp Cors")
    _seed_financial_periods(client, org_id)

    resp = client.get(
        f"/api/v1/organizations/{org_id}/exports/financial_periods",
        params={"format": "csv"},
        headers={"Origin": "http://localhost:3000"},
    )
    assert resp.status_code == 200
    exposed = resp.headers.get("access-control-expose-headers", "")
    assert "Content-Disposition" in exposed


def test_export_xlsx_is_a_valid_workbook(client):
    org_id = _register_login_and_create_org(client, "exp.xlsx@example.com", "Empresa Exp Xlsx")
    _seed_financial_periods(client, org_id)

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "xlsx"})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    sheet = wb.active
    assert sheet.max_row == 3  # cabecalho + 2 linhas
    assert sheet.cell(row=1, column=1).value == "Período"


def test_export_json_uses_canonical_keys_and_is_reimportable_shape(client):
    org_id = _register_login_and_create_org(client, "exp.json@example.com", "Empresa Exp Json")
    _seed_financial_periods(client, org_id)

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "json"})
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    assert payload[0]["period"] == "2026-01-01"
    assert payload[0]["revenue"] == 10000.0


def test_export_filters_by_date_range(client):
    org_id = _register_login_and_create_org(client, "exp.range@example.com", "Empresa Exp Range")
    _seed_financial_periods(client, org_id)

    resp = client.get(
        f"/api/v1/organizations/{org_id}/exports/financial_periods",
        params={"format": "json", "date_from": "2026-02-01", "date_to": "2026-02-28"},
    )
    payload = json.loads(resp.content)
    assert len(payload) == 1
    assert payload[0]["period"] == "2026-02-01"
    assert (
        'filename="financial_periods_2026-02-01_2026-02-28.json"' in resp.headers["content-disposition"]
    )


def test_export_empty_organization_returns_empty_list(client):
    org_id = _register_login_and_create_org(client, "exp.empty@example.com", "Empresa Exp Empty")
    resp = client.get(f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "json"})
    assert resp.status_code == 200
    assert json.loads(resp.content) == []


def test_export_unsupported_domain_is_rejected(client):
    org_id = _register_login_and_create_org(client, "exp.baddomain@example.com", "Empresa Exp Bad Domain")
    resp = client.get(f"/api/v1/organizations/{org_id}/exports/leads", params={"format": "csv"})
    assert resp.status_code == 400


def test_export_unsupported_format_is_rejected(client):
    org_id = _register_login_and_create_org(client, "exp.badformat@example.com", "Empresa Exp Bad Format")
    resp = client.get(
        f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "pdf"}
    )
    assert resp.status_code == 422  # Literal invalido -> validacao do FastAPI


def test_export_funnel_snapshots_uses_integer_values(client):
    org_id = _register_login_and_create_org(client, "exp.funnel@example.com", "Empresa Exp Funil")
    client.put(
        f"/api/v1/organizations/{org_id}/funnel-snapshots",
        json={"period": "2026-01-01", "impressions": 10000, "visits": 2000, "leads": 150},
    )

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/funnel_snapshots", params={"format": "json"})
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    assert payload[0]["impressions"] == 10000
    assert payload[0]["leads"] == 150
    assert payload[0]["mqls"] is None


def test_export_products_includes_inactive_and_ignores_date_range(client):
    org_id = _register_login_and_create_org(client, "exp.products@example.com", "Empresa Exp Produtos")
    client.post(f"/api/v1/organizations/{org_id}/products", json={"name": "Notebook Pro", "price": 5000})
    inactive = client.post(
        f"/api/v1/organizations/{org_id}/products", json={"name": "Produto Descontinuado", "price": 10}
    ).json()
    client.patch(f"/api/v1/organizations/{org_id}/products/{inactive['id']}", json={"active": False})

    # date_from/date_to são ignorados — Produtos não tem conceito de período
    # (DomainSpec.supports_date_range = False) — e mesmo assim devolve tudo,
    # inclusive o produto inativo (RN-014: nunca excluído de verdade).
    resp = client.get(
        f"/api/v1/organizations/{org_id}/exports/products",
        params={"format": "json", "date_from": "2020-01-01", "date_to": "2020-12-31"},
    )
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    names = {p["name"] for p in payload}
    assert names == {"Notebook Pro", "Produto Descontinuado"}
    by_name = {p["name"]: p for p in payload}
    assert by_name["Notebook Pro"]["active"] is True
    assert by_name["Produto Descontinuado"]["active"] is False
    # Filename não menciona um intervalo que não foi de fato aplicado.
    assert 'filename="products_export.json"' in resp.headers["content-disposition"]


def test_export_campaigns_resolves_channel_name(client):
    org_id = _register_login_and_create_org(client, "exp.campaigns@example.com", "Empresa Exp Campanhas")
    channel = client.post(
        f"/api/v1/organizations/{org_id}/marketing-channels", json={"name": "Google Ads"}
    ).json()
    client.post(
        f"/api/v1/organizations/{org_id}/campaigns",
        json={"channel_id": channel["id"], "name": "Black Friday", "investment": 5000, "clicks": 1200},
    )

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/campaigns", params={"format": "json"})
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    # "channel" sai como nome, nunca o channel_id interno.
    assert payload[0]["channel"] == "Google Ads"
    assert payload[0]["name"] == "Black Friday"
    assert payload[0]["investment"] == 5000.0
    assert payload[0]["clicks"] == 1200


def test_export_audit_logs_contains_seeded_entries_with_resolved_user_email(client):
    org_id = _register_login_and_create_org(
        client, "exp.auditcontent@example.com", "Empresa Exp Audit Content"
    )
    _seed_financial_periods(client, org_id)  # PUT /financials grava um audit_log

    resp = client.get(f"/api/v1/organizations/{org_id}/exports/audit_logs", params={"format": "json"})
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    upserts = [row for row in payload if row["action"] == "financial_period.upserted"]
    assert len(upserts) == 2  # um por PUT /financials em _seed_financial_periods
    assert upserts[0]["user_email"] == "exp.auditcontent@example.com"
    assert upserts[0]["entity"] == "financial_period"
    assert upserts[0]["entity_id"] is not None


def test_export_audit_logs_filters_by_date_range(client):
    org_id = _register_login_and_create_org(client, "exp.auditrange@example.com", "Empresa Exp Audit Range")
    _seed_financial_periods(client, org_id)

    # "Hoje" ancorado em UTC, igual ao filtro de produção
    # (audit/repository.py::list_audit_logs_in_range filtra por dia de calendário
    # inteiro em UTC porque CreatedAtMixin usa func.clock_timestamp()).
    today = datetime.now(UTC).date().isoformat()
    resp_today = client.get(
        f"/api/v1/organizations/{org_id}/exports/audit_logs",
        params={"format": "json", "date_from": today, "date_to": today},
    )
    assert len(json.loads(resp_today.content)) >= 2

    resp_past = client.get(
        f"/api/v1/organizations/{org_id}/exports/audit_logs",
        params={"format": "json", "date_from": "2020-01-01", "date_to": "2020-01-02"},
    )
    assert json.loads(resp_past.content) == []


def test_export_audit_logs_isolates_organizations(client):
    org_a = _register_login_and_create_org(client, "exp.auditisoa@example.com", "Empresa Exp Audit Iso A")
    _seed_financial_periods(client, org_a)
    org_b = client.post("/api/v1/organizations", json={"name": "Empresa Exp Audit Iso B"}).json()["id"]

    resp = client.get(f"/api/v1/organizations/{org_b}/exports/audit_logs", params={"format": "json"})
    assert resp.status_code == 200
    payload = json.loads(resp.content)
    assert all(row["action"] != "financial_period.upserted" for row in payload)


def test_viewer_cannot_export_audit_logs(client_factory):
    owner = client_factory()
    viewer_client = client_factory()
    _register_and_login(owner, "Ana", "exp.auditviewerowner@example.com")
    _register_and_login(viewer_client, "Vitor", "exp.auditviewer@example.com")

    org_id = owner.post("/api/v1/organizations", json={"name": "Empresa Exp Audit Viewer"}).json()["id"]
    owner.post(
        f"/api/v1/organizations/{org_id}/memberships",
        json={"email": "exp.auditviewer@example.com", "role": "VIEWER"},
    )

    resp = viewer_client.get(f"/api/v1/organizations/{org_id}/exports/audit_logs", params={"format": "json"})
    assert resp.status_code == 403


def test_analyst_cannot_export_audit_logs(client_factory):
    owner = client_factory()
    analyst_client = client_factory()
    _register_and_login(owner, "Ana", "exp.auditanalystowner@example.com")
    _register_and_login(analyst_client, "Beto", "exp.auditanalyst@example.com")

    org_id = owner.post("/api/v1/organizations", json={"name": "Empresa Exp Audit Analyst"}).json()["id"]
    owner.post(
        f"/api/v1/organizations/{org_id}/memberships",
        json={"email": "exp.auditanalyst@example.com", "role": "ANALYST"},
    )

    resp = analyst_client.get(f"/api/v1/organizations/{org_id}/exports/audit_logs", params={"format": "json"})
    assert resp.status_code == 403


def test_organization_admin_can_export_audit_logs(client_factory):
    owner = client_factory()
    admin_client = client_factory()
    _register_and_login(owner, "Ana", "exp.auditadminowner@example.com")
    _register_and_login(admin_client, "Carla", "exp.auditadmin@example.com")

    org_id = owner.post("/api/v1/organizations", json={"name": "Empresa Exp Audit Admin"}).json()["id"]
    owner.post(
        f"/api/v1/organizations/{org_id}/memberships",
        json={"email": "exp.auditadmin@example.com", "role": "ORGANIZATION_ADMIN"},
    )

    resp = admin_client.get(f"/api/v1/organizations/{org_id}/exports/audit_logs", params={"format": "json"})
    assert resp.status_code == 200


def test_viewer_can_export_same_as_dashboard_read(client_factory):
    owner = client_factory()
    viewer_client = client_factory()
    _register_and_login(owner, "Ana", "exp.viewerowner@example.com")
    _register_and_login(viewer_client, "Vitor", "exp.viewer@example.com")

    org_id = owner.post("/api/v1/organizations", json={"name": "Empresa Exp Viewer"}).json()["id"]
    owner.post(
        f"/api/v1/organizations/{org_id}/memberships",
        json={"email": "exp.viewer@example.com", "role": "VIEWER"},
    )
    _seed_financial_periods(owner, org_id)

    resp = viewer_client.get(
        f"/api/v1/organizations/{org_id}/exports/financial_periods", params={"format": "csv"}
    )
    assert resp.status_code == 200
